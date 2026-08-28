"""Normalize a wide rolling-timetable workbook into stored week rows.

The parser in ``rolling_timetable.py`` is for unit-delivery analysis. This
module reads the same workbook for persistence: Intake labels come from the
header row, NA cells are dropped, and every other cell becomes one row.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from typing import Iterable, Sequence

from app.services.rolling_timetable import (
    ASSESSMENT,
    BREAK,
    INTAKE_MARKER,
    RollingTimetableError,
    _as_date,
    _clean,
    _locate_columns,
)

TYPE_UNIT = "UNIT"
TYPE_BREAK = "BREAK"
TYPE_ASSESSMENT = "ASSESSMENT_WEEK"
TYPE_NA = "NA"
TYPE_INTAKE = "INTAKE"

DISPLAY_BREAK = "Break"
DISPLAY_ASSESSMENT = "Assessment Week"
DISPLAY_NA = "NA"

SHEET_NAME_RE = re.compile(r"^(?P<code>.+)_(?P<duration>\d+)_Weeks$")
INTAKE_LABEL_RE = re.compile(
    r"^(?P<code>.+)_(?P<duration>\d+)_(?P<date>.+)_(?P<group>.+)_Intake$"
)
INTAKE_DATE_FORMAT = "%d %b %Y"


@dataclass
class NormalizedWeekRow:
    qualification_code: str
    duration_weeks: int
    intake_label: str
    intake_group: str
    intake_start_date: dt.date
    week_no: int
    week_start_date: dt.date
    week_end_date: dt.date
    schedule_type: str
    schedule_value: str
    unit_code: str | None
    unit_code_2: str | None
    unit_count: int
    unit_delivery_span_weeks: int | None = None


@dataclass
class NormalizeReport:
    sheets: int = 0
    intakes: int = 0
    week_rows_read: int = 0
    cells_read: int = 0
    stored_rows: int = 0
    unit_cells: int = 0
    break_cells: int = 0
    assessment_cells: int = 0
    na_cells: int = 0
    skipped_intake_markers: int = 0
    unknown_values: list[str] = field(default_factory=list)

    def as_lines(self) -> list[str]:
        return [
            f"  sheets              : {self.sheets}",
            f"  intakes             : {self.intakes}",
            f"  week rows read      : {self.week_rows_read}",
            f"  cells read          : {self.cells_read}",
            f"  stored rows         : {self.stored_rows}",
            f"  UNIT                : {self.unit_cells}",
            f"  Break               : {self.break_cells}",
            f"  Assessment Week     : {self.assessment_cells}",
            f"  NA (not stored)     : {self.na_cells}",
            f"  Intake markers skip : {self.skipped_intake_markers}",
        ]


def parse_sheet_identity(sheet_name: str) -> tuple[str, int]:
    match = SHEET_NAME_RE.match(sheet_name.strip())
    if not match:
        raise RollingTimetableError(
            f"Sheet {sheet_name!r} does not match QualificationCode_Duration_Weeks."
        )
    return match.group("code").upper(), int(match.group("duration"))


def parse_intake_label(label: str) -> tuple[str, int, dt.date, str]:
    match = INTAKE_LABEL_RE.match(label.strip())
    if not match:
        raise RollingTimetableError(
            f"Intake label {label!r} does not match "
            "Qualification_Duration_Date_Group_Intake."
        )
    try:
        start = dt.datetime.strptime(match.group("date"), INTAKE_DATE_FORMAT).date()
    except ValueError as exc:
        raise RollingTimetableError(
            f"Intake label {label!r} has an unreadable date {match.group('date')!r}."
        ) from exc
    return (
        match.group("code").upper(),
        int(match.group("duration")),
        start,
        match.group("group"),
    )


def classify_cell(value: str) -> tuple[str, str]:
    """Return (schedule_type, stored_display_value)."""
    text = value.strip()
    upper = text.upper()
    if upper == TYPE_NA or upper == DISPLAY_NA:
        return TYPE_NA, DISPLAY_NA
    if upper == BREAK:
        return TYPE_BREAK, text if text != upper else DISPLAY_BREAK
    if upper == ASSESSMENT:
        return TYPE_ASSESSMENT, text if text != upper else DISPLAY_ASSESSMENT
    if upper == INTAKE_MARKER:
        return TYPE_INTAKE, text
    return TYPE_UNIT, text


def split_unit_codes(schedule_value: str) -> tuple[str, str | None, int]:
    parts = [part.strip() for part in schedule_value.split("/") if part.strip()]
    if not parts:
        raise RollingTimetableError(f"Unit cell {schedule_value!r} split to nothing.")
    if len(parts) == 1:
        return parts[0], None, 1
    if len(parts) == 2:
        return parts[0], parts[1], 2
    # More than two codes: keep the source value, store first two splits.
    return parts[0], parts[1], len(parts)


def normalize_sheet(
    sheet_name: str,
    rows: Sequence[Sequence[object]],
) -> tuple[list[NormalizedWeekRow], NormalizeReport]:
    qualification_code, duration_weeks = parse_sheet_identity(sheet_name)
    if not rows:
        raise RollingTimetableError(f"Sheet {sheet_name!r} is empty.")

    header = list(rows[0])
    week_col, start_col, end_col = _locate_columns(header)
    fixed = {week_col, start_col, end_col}

    intakes: list[tuple[int, str, str, dt.date, str]] = []
    for index, cell in enumerate(header):
        if index in fixed:
            continue
        label = _clean(cell)
        if not label:
            continue
        code, duration, start, group = parse_intake_label(label)
        if code != qualification_code or duration != duration_weeks:
            raise RollingTimetableError(
                f"Sheet {sheet_name!r}: intake {label!r} does not match the sheet identity."
            )
        intakes.append((index, label, group, start, code))

    if not intakes:
        raise RollingTimetableError(f"Sheet {sheet_name!r} has no intake columns.")

    report = NormalizeReport(sheets=1, intakes=len(intakes))
    raw_rows: list[NormalizedWeekRow] = []

    for line_no, row in enumerate(rows[1:], start=2):
        if all(_clean(c) == "" for c in row):
            continue
        raw_week = row[week_col] if week_col < len(row) else None
        start = _as_date(row[start_col]) if start_col < len(row) else None
        end = _as_date(row[end_col]) if end_col < len(row) else None
        if raw_week is None or start is None or end is None:
            continue
        try:
            week_no = int(raw_week)
        except (TypeError, ValueError):
            raise RollingTimetableError(
                f"{sheet_name} row {line_no}: Week No. {raw_week!r} is not a whole number."
            ) from None
        if end < start:
            raise RollingTimetableError(
                f"{sheet_name} row {line_no}: End Date {end} is before Start Date {start}."
            )
        report.week_rows_read += 1

        for column, label, group, intake_start, _code in intakes:
            cell = row[column] if column < len(row) else None
            value = _clean(cell)
            report.cells_read += 1
            if not value:
                # Spec: the approved workbook has no blanks. Treat blank as NA
                # so a missing cell cannot become a phantom unit.
                report.na_cells += 1
                continue
            schedule_type, schedule_value = classify_cell(value)
            if schedule_type == TYPE_NA:
                report.na_cells += 1
                continue
            if schedule_type == TYPE_INTAKE:
                report.skipped_intake_markers += 1
                continue
            unit_code = unit_code_2 = None
            unit_count = 0
            if schedule_type == TYPE_UNIT:
                unit_code, unit_code_2, unit_count = split_unit_codes(schedule_value)
                report.unit_cells += 1
            elif schedule_type == TYPE_BREAK:
                report.break_cells += 1
            else:
                report.assessment_cells += 1

            raw_rows.append(
                NormalizedWeekRow(
                    qualification_code=qualification_code,
                    duration_weeks=duration_weeks,
                    intake_label=label,
                    intake_group=group,
                    intake_start_date=intake_start,
                    week_no=week_no,
                    week_start_date=start,
                    week_end_date=end,
                    schedule_type=schedule_type,
                    schedule_value=schedule_value,
                    unit_code=unit_code,
                    unit_code_2=unit_code_2,
                    unit_count=unit_count,
                )
            )

    _stamp_delivery_spans(raw_rows)
    raw_rows, duplicate_conflicts = _dedupe_identical_weeks(raw_rows)
    if duplicate_conflicts:
        raise RollingTimetableError(
            f"{sheet_name}: conflicting duplicate week rows: " + "; ".join(duplicate_conflicts[:8])
        )
    report.stored_rows = len(raw_rows)
    return raw_rows, report


def _dedupe_identical_weeks(
    rows: list[NormalizedWeekRow],
) -> tuple[list[NormalizedWeekRow], list[str]]:
    """Keep one row per business key when the source repeated an identical week.

    BSB60720 repeats weeks 17-20 with the same dates and Break values. That is
    not two requirements — it is a duplicated calendar block. Conflicting
    duplicates (same week, different value) are refused.
    """
    kept: dict[tuple[str, int, str, int], NormalizedWeekRow] = {}
    conflicts: list[str] = []
    for row in rows:
        key = (row.qualification_code, row.duration_weeks, row.intake_label, row.week_no)
        existing = kept.get(key)
        if existing is None:
            kept[key] = row
            continue
        same = (
            existing.week_start_date == row.week_start_date
            and existing.week_end_date == row.week_end_date
            and existing.schedule_value == row.schedule_value
            and existing.schedule_type == row.schedule_type
        )
        if not same:
            conflicts.append(
                f"{row.intake_label} week {row.week_no}: "
                f"{existing.schedule_value!r} vs {row.schedule_value!r}"
            )
    return list(kept.values()), conflicts


def _stamp_delivery_spans(rows: list[NormalizedWeekRow]) -> None:
    """Calendar span of one continuous unit delivery, breaks included.

    A break or assessment week does not split a unit. A different schedule
    value does. Matches the approved parser behaviour.
    """
    by_intake: dict[tuple[str, int, str], list[NormalizedWeekRow]] = {}
    for row in rows:
        by_intake.setdefault(
            (row.qualification_code, row.duration_weeks, row.intake_label), []
        ).append(row)

    for group in by_intake.values():
        group.sort(key=lambda row: row.week_no)
        pending_value: str | None = None
        pending_rows: list[NormalizedWeekRow] = []
        pending_start: dt.date | None = None
        pending_end: dt.date | None = None
        gap = False

        def close() -> None:
            nonlocal pending_value, pending_rows, pending_start, pending_end, gap
            if pending_value is None or pending_start is None or pending_end is None:
                pending_value = None
                pending_rows = []
                return
            span = ((pending_end - pending_start).days + 1) // 7
            for item in pending_rows:
                item.unit_delivery_span_weeks = span
            pending_value = None
            pending_rows = []
            pending_start = None
            pending_end = None
            gap = False

        for row in group:
            if row.schedule_type == TYPE_UNIT:
                if pending_value is None:
                    pending_value = row.schedule_value
                    pending_rows = [row]
                    pending_start = row.week_start_date
                    pending_end = row.week_end_date
                    gap = False
                elif row.schedule_value == pending_value:
                    pending_rows.append(row)
                    pending_end = row.week_end_date
                    gap = False
                else:
                    close()
                    pending_value = row.schedule_value
                    pending_rows = [row]
                    pending_start = row.week_start_date
                    pending_end = row.week_end_date
                    gap = False
            else:
                # Break / assessment: hold the delivery open until a different unit.
                gap = True
        close()


def merge_reports(reports: Iterable[NormalizeReport]) -> NormalizeReport:
    merged = NormalizeReport()
    for report in reports:
        merged.sheets += report.sheets
        merged.intakes += report.intakes
        merged.week_rows_read += report.week_rows_read
        merged.cells_read += report.cells_read
        merged.stored_rows += report.stored_rows
        merged.unit_cells += report.unit_cells
        merged.break_cells += report.break_cells
        merged.assessment_cells += report.assessment_cells
        merged.na_cells += report.na_cells
        merged.skipped_intake_markers += report.skipped_intake_markers
        merged.unknown_values.extend(report.unknown_values)
    return merged


def load_workbook_rows(path: str, *, qualification: str | None = None) -> tuple[list[NormalizedWeekRow], NormalizeReport]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        all_rows: list[NormalizedWeekRow] = []
        reports: list[NormalizeReport] = []
        for sheet_name in workbook.sheetnames:
            if not SHEET_NAME_RE.match(sheet_name):
                continue
            code, _duration = parse_sheet_identity(sheet_name)
            if qualification and code.upper() != qualification.upper():
                continue
            sheet = workbook[sheet_name]
            rows, report = normalize_sheet(sheet_name, list(sheet.iter_rows(values_only=True)))
            all_rows.extend(rows)
            reports.append(report)
        if not reports:
            raise RollingTimetableError(
                f"No rolling timetable sheets found in {path}."
                + (f" Qualification filter: {qualification}." if qualification else "")
            )
        return all_rows, merge_reports(reports)
    finally:
        workbook.close()
