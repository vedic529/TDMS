"""Rebuild the uploaded column layout from stored allocation tables."""

from __future__ import annotations

import csv
import datetime as dt
import io
from collections import defaultdict

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.allocation import AllocationDelivery, AllocationSession
from app.models.college import Campus, College
from app.models.facility import Facility
from app.models.qualification import Qualification, Unit
from app.models.trainer import Trainer
from app.services.allocation_profiles import ALLOCATION_COLUMNS

#: The spreadsheet download is one file, not a paged view.
_EXPORT_LIMIT = 100_000

WEEKDAY_LABEL = {
    "MONDAY": "Monday",
    "TUESDAY": "Tuesday",
    "WEDNESDAY": "Wednesday",
    "THURSDAY": "Thursday",
    "FRIDAY": "Friday",
    "SATURDAY": "Saturday",
}


def _format_time(value: dt.time) -> str:
    hour = value.hour % 12 or 12
    meridian = "am" if value.hour < 12 else "pm"
    if value.minute:
        return f"{hour}:{value.minute:02d} {meridian}"
    return f"{hour} {meridian}"


def _days_and_times(sessions: list[AllocationSession]) -> str:
    lines = [
        f"{WEEKDAY_LABEL[row.weekday]}- {_format_time(row.start_time)} to {_format_time(row.end_time)}"
        for row in sorted(sessions, key=lambda item: item.weekday)
    ]
    return "\n".join(lines)


def _shared_or_per_day(sessions: list[AllocationSession], getter) -> str:
    values = [(row.weekday, getter(row) or "") for row in sorted(sessions, key=lambda item: item.weekday)]
    unique = {value for _, value in values if value}
    if len(unique) <= 1:
        return next(iter(unique), "NA")
    return "\n".join(f"{WEEKDAY_LABEL[day]}- {value}" for day, value in values if value)


def _approved(lookup: dict, key: int | None, attribute: str, fallback: str | None) -> str:
    """The approved value, else the text the file supplied, else blank.

    Every one of these foreign keys became nullable in `c8f3b2a1d470`, so a
    direct `lookup[key]` raised `KeyError: None` and one unresolved row crashed
    the whole export. Falling back to the stored text shows what the file
    actually said rather than a blank.
    """
    if key is not None:
        record = lookup.get(key)
        if record is not None:
            return getattr(record, attribute, None) or ""
    return fallback or ""


def export_workbook(session: Session, *, training_package: str, start_date: dt.date, end_date: dt.date | None) -> tuple[str, bytes]:
    package = training_package.upper()
    conditions = [
        AllocationDelivery.training_package == package,
        AllocationDelivery.end_date >= start_date,
        AllocationDelivery.is_quarantined.is_(False),
    ]
    if end_date is not None:
        conditions.append(AllocationDelivery.start_date <= end_date)
    deliveries = list(
        session.execute(
            select(AllocationDelivery)
            .where(*conditions)
            .order_by(AllocationDelivery.start_date, AllocationDelivery.id)
        ).scalars()
    )
    if not deliveries:
        return _xlsx([], [])
    sessions = list(
        session.execute(
            select(AllocationSession).where(
                AllocationSession.training_package == package,
                AllocationSession.delivery_id.in_([row.id for row in deliveries]),
            )
        ).scalars()
    )
    by_delivery: dict[int, list[AllocationSession]] = defaultdict(list)
    for row in sessions:
        by_delivery[row.delivery_id].append(row)
    colleges = {row.id: row for row in session.execute(select(College)).scalars()}
    campuses = {row.id: row for row in session.execute(select(Campus)).scalars()}
    quals = {row.id: row for row in session.execute(select(Qualification)).scalars()}
    units = {row.id: row for row in session.execute(select(Unit)).scalars()}
    facilities = {row.id: row for row in session.execute(select(Facility)).scalars()}
    trainers = {row.id: row for row in session.execute(select(Trainer)).scalars()}

    output_rows: list[list[str]] = []
    for delivery in deliveries:
        grouped = defaultdict(list)
        for row in by_delivery.get(delivery.id, []):
            grouped[row.stream].append(row)
        theory = grouped.get("THEORY", [])
        practical = grouped.get("PRACTICAL", [])
        mscris = grouped.get("MSCRIS", [])
        def room(row: AllocationSession) -> str:
            if row.virtual_kind == "FACE_TO_FACE_VC":
                return "Face to Face VC"
            if row.facility_id and row.facility_id in facilities:
                return facilities[row.facility_id].facility_reference
            return row.classroom_text or ""

        def trainer(row: AllocationSession) -> str:
            if row.trainer_id and row.trainer_id in trainers:
                return trainers[row.trainer_id].trainer_name
            return row.trainer_text or ""

        output_rows.append(
            [
                _approved(colleges, delivery.college_id, "college_full_name", delivery.college_text),
                _approved(campuses, delivery.campus_id, "campus_location", delivery.campus_text),
                _approved(quals, delivery.qualification_id, "qualification_code", delivery.qualification_text),
                _approved(quals, delivery.qualification_id, "qualification_title", ""),
                str(delivery.duration_weeks),
                delivery.group_code,
                "" if delivery.classroom_size is None else str(delivery.classroom_size),
                _approved(units, delivery.unit_id, "unit_code", delivery.unit_text),
                _approved(units, delivery.unit_id, "unit_title", ""),
                delivery.uoc_type.replace("_", " ").title().replace("And", "and") if delivery.uoc_type != "THEORY_AND_PRACTICAL" else "Theory and Practical",
                delivery.mode_of_delivery,
                delivery.start_date.isoformat(),
                delivery.end_date.isoformat(),
                _days_and_times(theory) or "NA",
                _shared_or_per_day(theory, room) if theory else "NA",
                "NA",
                _shared_or_per_day(theory, trainer) if theory else "NA",
                "NA",
                "NA",
                "NA",
                "NA",
                _shared_or_per_day(mscris, room) if mscris else "NA",
                _days_and_times(mscris) or "NA",
                _shared_or_per_day(mscris, trainer) if mscris else "NA",
                delivery.remarks or "NA",
            ]
        )
        # keep practical columns NA for BSB; include them when present
        if practical:
            output_rows[-1][17] = _shared_or_per_day(practical, room)
            output_rows[-1][19] = _days_and_times(practical)
            output_rows[-1][20] = _shared_or_per_day(practical, trainer)
    return _xlsx(list(ALLOCATION_COLUMNS), output_rows)


def _xlsx(headers: list[str], rows: list[list[str]]) -> tuple[str, bytes]:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Allocation"
    sheet.append(list(ALLOCATION_COLUMNS if not headers else headers))
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return "allocation.xlsx", buffer.getvalue()


# ---------------------------------------------------------------------------
# The spreadsheet view, downloaded (approved 21 September 2026)
# ---------------------------------------------------------------------------

#: The spreadsheet view's columns, in its order. One list, so the screen and the
#: download cannot drift apart.
SPREADSHEET_COLUMNS = (
    "Sl No.",
    "College",
    "Campus",
    "Qualification Code",
    "Qualification Title",
    "Duration in weeks",
    "Group",
    "Linked Intakes",
    "Total students attending",
    "COE students",
    "Non-COE students",
    "Units of Competency ID",
    "Units of Competency Title",
    "Unit of Competency Start Date",
    "Unit of Competency End Date",
    "UoC Type",
    "Mode of Delivery",
    "Theory Class Days and Times",
    "Theory Classroom Name",
    "Theory Classroom Capacity",
    "Theory Trainer",
    "Practical Classroom Name",
    "Practical Class Capacity",
    "Practical Class Days and Times",
    "Practical Trainers",
)

EMPTY = "—"


def _friendly(value: str) -> str:
    """`THEORY_ONLY` reads as `Theory Only`, exactly as the screen shows it."""
    return " ".join(part.capitalize() for part in str(value or "").split("_"))


def _stacked(values: list[str], empty: str = EMPTY) -> str:
    """Several values in one cell, one per line — the screen stacks them."""
    seen = list(dict.fromkeys(values))
    return "\n".join(seen) if seen else empty


def _session_values(items: list[dict], field: str, empty: str = EMPTY) -> str:
    values = []
    for item in items:
        if field == "schedule":
            values.append(f"{_friendly(item['weekday'])} · {item['start_time']}–{item['end_time']}")
        elif field == "classroom":
            values.append(item["classroom"] or EMPTY)
        elif field == "capacity":
            capacity = item["classroom_capacity"]
            values.append(
                str(capacity) if capacity is not None else ("Virtual" if item["delivery_mode"] == "VIRTUAL" else EMPTY)
            )
        else:
            values.append(item["trainer"] or EMPTY)
    return _stacked(values, empty)


def export_spreadsheet(
    session: Session,
    *,
    training_package: str,
    start_date: dt.date,
    end_date: dt.date | None,
    file_format: str = "xlsx",
) -> tuple[str, bytes, str]:
    """The allocation spreadsheet exactly as the screen shows it.

    Reads the same service the view reads, so the columns, the order and the
    wording are the screen's - not a second layout that drifts from it. Returns
    (file name, bytes, media type).
    """
    from app.services import allocation_spreadsheet

    package = training_package.upper()
    page = allocation_spreadsheet.list_rows(
        session,
        training_package=package,
        start_date=start_date,
        end_date=end_date,
        limit=_EXPORT_LIMIT,
        offset=0,
    )

    rows: list[list[str]] = []
    for item in page["items"]:
        theory = [row for row in item["sessions"] if row["stream"] == "THEORY"]
        practical = [row for row in item["sessions"] if row["stream"] == "PRACTICAL"]
        rows.append(
            [
                str(item["sl_no"]),
                item["college"] or EMPTY,
                item["campus"] or EMPTY,
                item["qualification_code"] or EMPTY,
                item["qualification_title"] or EMPTY,
                str(item["duration_weeks"]),
                item["group"] or EMPTY,
                _stacked(list(item["intakes"])),
                str(item["total_students"]),
                str(item["coe_students"]),
                str(item["non_coe_students"]),
                item["unit_code"] or EMPTY,
                item["unit_title"] or EMPTY,
                item["unit_start_date"].isoformat(),
                item["unit_end_date"].isoformat(),
                _friendly(item["uoc_type"]),
                item["mode_of_delivery"] or EMPTY,
                _session_values(theory, "schedule"),
                _session_values(theory, "classroom"),
                _session_values(theory, "capacity"),
                _session_values(theory, "trainer"),
                _session_values(practical, "classroom"),
                _session_values(practical, "capacity"),
                _session_values(
                    practical, "schedule", "Not required" if item["uoc_type"] == "THEORY_ONLY" else EMPTY
                ),
                _session_values(practical, "trainer"),
            ]
        )

    stamp = dt.date.today().isoformat()
    if file_format == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(SPREADSHEET_COLUMNS)
        writer.writerows(rows)
        return (
            f"allocation-records-{package}-{stamp}.csv",
            buffer.getvalue().encode("utf-8-sig"),
            "text/csv",
        )

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Allocation Records"
    sheet.append(list(SPREADSHEET_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    # The screen stacks several values in a cell; the workbook must show them
    # the same way rather than running them together on one line.
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    output = io.BytesIO()
    book.save(output)
    return (
        f"allocation-records-{package}-{stamp}.xlsx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
